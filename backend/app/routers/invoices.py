import hashlib
import logging
from typing import Optional, List
from io import BytesIO
from datetime import date
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException, Query, BackgroundTasks, Request
from fastapi.responses import Response, StreamingResponse
from fastapi.concurrency import run_in_threadpool
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
from decimal import Decimal

from app.database import get_db
from app.models.invoice import Invoice, OcrResult, LlmResult, ParsingDiff, InvoiceStatus
from app.schemas.invoice import (
    InvoiceResponse, InvoiceListResponse, InvoiceDetailResponse,
    InvoiceUpdate, BatchUpdateRequest, BatchDeleteRequest, StatisticsResponse, UploadResponse,
    FieldUpdateRequest, QrRetryResponse
)
from app.config import get_settings
from app.services.audit_service import log_audit_no_commit, get_client_info
from app.services import app_settings
from app.services.fusion import AMOUNT_GROUP, PIPELINE_VERSION, SRC_MANUAL
from app.services.review import (
    DIFF_FIELDS, REQUIRED_FIELDS, InvalidValue,
    apply_to_invoice, reconcile_after_edit, resolve_set, restore_auto, revalidate,
)
from app.services.validation import VERDICT_REJECTED
from app.rate_limit import limiter

settings = get_settings()
router = APIRouter()
logger = logging.getLogger(__name__)

# 已进入报销流程的发票只读
LOCKED_STATUSES = {InvoiceStatus.REIMBURSED, InvoiceStatus.NOT_REIMBURSED}
# 还没有识别结果的发票，没有可以核对的字段
UNPROCESSED_STATUSES = {InvoiceStatus.UPLOADED, InvoiceStatus.PROCESSING}

# 预览渲染方式变了（DPI、格式）要改这个版本号，否则浏览器会继续用缓存
PREVIEW_VERSION = "v1"


def _require_editable(invoice: Invoice, allow_unprocessed: bool = False) -> None:
    if invoice.status in LOCKED_STATUSES:
        raise HTTPException(status_code=409, detail="发票已进入报销流程，不能修改")
    if not allow_unprocessed and invoice.status in UNPROCESSED_STATUSES:
        raise HTTPException(status_code=409, detail="发票还没有识别完成，暂时不能修改")


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


async def _get_invoice_or_404(db: AsyncSession, invoice_id: int, for_update: bool = False) -> Invoice:
    query = select(Invoice).where(Invoice.id == invoice_id)
    if for_update:
        # 同一张票的字段操作要串行：每次都会读-改-写 checks / verdict 与派生金额
        query = query.with_for_update()
    invoice = (await db.execute(query)).scalar_one_or_none()
    if not invoice:
        raise HTTPException(status_code=404, detail="发票不存在")
    return invoice


async def _get_diffs(db: AsyncSession, invoice_id: int) -> List[ParsingDiff]:
    query = (select(ParsingDiff)
             .where(ParsingDiff.invoice_id == invoice_id)
             .order_by(ParsingDiff.id))
    return list((await db.execute(query)).scalars().all())


def _refresh_verdict(invoice: Invoice, diffs: List[ParsingDiff]) -> None:
    """人工改动后按当前字段值重校验，回写 checks / verdict。

    v6 之前的旧票没有置信信息，规则口径也不同，保持原样。
    """
    if invoice.pipeline_version != PIPELINE_VERSION or not diffs:
        return
    cfg = app_settings.current()
    checks, verdict = revalidate(diffs, invoice.nulled or [], cfg.entities, cfg.amount_rules)
    invoice.checks = [c.as_dict() for c in checks]
    invoice.verdict = verdict


async def _build_detail(db: AsyncSession, invoice: Invoice) -> InvoiceDetailResponse:
    ocr = (await db.execute(
        select(OcrResult).where(OcrResult.invoice_id == invoice.id))).scalar_one_or_none()
    llm = (await db.execute(
        select(LlmResult).where(LlmResult.invoice_id == invoice.id))).scalar_one_or_none()
    diffs = await _get_diffs(db, invoice.id)

    # Build response dict from invoice to avoid lazy loading issues
    invoice_dict = {
        "id": invoice.id,
        "file_name": invoice.file_name,
        "file_type": invoice.file_type,
        "status": invoice.status,
        "owner": invoice.owner,
        "verdict": invoice.verdict,
        "checks": invoice.checks,
        "degraded": invoice.degraded,
        "nulled": invoice.nulled,
        "pipeline_version": invoice.pipeline_version,
        "invoice_code": invoice.invoice_code,
        "invoice_number": invoice.invoice_number,
        "issue_date": invoice.issue_date,
        "buyer_name": invoice.buyer_name,
        "buyer_tax_id": invoice.buyer_tax_id,
        "seller_name": invoice.seller_name,
        "seller_tax_id": invoice.seller_tax_id,
        "item_name": invoice.item_name,
        "total_with_tax": invoice.total_with_tax,
        "specification": invoice.specification,
        "unit": invoice.unit,
        "quantity": invoice.quantity,
        "unit_price": invoice.unit_price,
        "amount": invoice.amount,
        "tax_rate": invoice.tax_rate,
        "tax_amount": invoice.tax_amount,
        "created_at": invoice.created_at,
        "updated_at": invoice.updated_at,
        "ocr_result": ocr,
        "llm_result": llm,
        "parsing_diffs": diffs,
    }
    return InvoiceDetailResponse.model_validate(invoice_dict)


def _parse_invoice_ids(invoice_ids: Optional[str]) -> Optional[List[int]]:
    if not invoice_ids:
        return None
    ids: List[int] = []
    for raw_id in invoice_ids.split(","):
        raw_id = raw_id.strip()
        if not raw_id:
            continue
        try:
            ids.append(int(raw_id))
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="发票ID必须为整数") from exc
    if not ids:
        raise HTTPException(status_code=400, detail="发票ID不能为空")
    return ids


def _parse_date_param(value: Optional[str], field_name: str) -> Optional[date]:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"{field_name} 日期格式无效，应为 YYYY-MM-DD") from exc


@router.post("/upload", response_model=List[UploadResponse])
@limiter.limit("10/minute")
async def upload_invoices(
    request: Request,
    background_tasks: BackgroundTasks,
    files: List[UploadFile] = File(...),
    db: AsyncSession = Depends(get_db)
):
    """上传发票文件 (支持多文件)，上传后异步触发OCR解析"""
    results = []
    invoice_ids_to_process = []
    client_info = get_client_info(request)

    for file in files:
        # Validate file type
        ext = file.filename.split(".")[-1].lower() if file.filename else ""
        if ext not in settings.allowed_extensions:
            results.append(UploadResponse(
                id=0,
                file_name=file.filename or "unknown",
                status="error",
                message=f"不支持的文件类型: {ext}"
            ))
            continue

        # Read file content
        content = await file.read()

        # Validate file size
        if len(content) > settings.max_file_size:
            results.append(UploadResponse(
                id=0,
                file_name=file.filename or "unknown",
                status="error",
                message=f"文件过大，最大支持 {settings.max_file_size // 1024 // 1024}MB"
            ))
            continue

        # Create invoice record with UPLOADED status (not yet processed)
        invoice = Invoice(
            file_name=file.filename or "unknown",
            file_type=ext,
            file_data=content,
            status=InvoiceStatus.UPLOADED
        )
        db.add(invoice)
        await db.flush()

        invoice_ids_to_process.append(invoice.id)

        # Audit log for upload
        await log_audit_no_commit(
            db=db,
            entity_type="invoice",
            entity_id=invoice.id,
            action="upload",
            new_value={"file_name": invoice.file_name, "file_type": ext, "file_size": len(content)},
            ip_address=client_info.get("ip_address"),
            user_agent=client_info.get("user_agent"),
        )

        results.append(UploadResponse(
            id=invoice.id,
            file_name=invoice.file_name,
            status="success",
            message="上传成功，等待解析"
        ))

    await db.commit()

    # Schedule background processing for each uploaded invoice
    for invoice_id in invoice_ids_to_process:
        background_tasks.add_task(process_invoice_background, invoice_id)

    return results


async def process_invoice_background(invoice_id: int, max_retries: int = 3):
    """Background task to process an invoice with OCR/LLM.

    Implements exponential backoff retry logic for transient failures.
    Total attempts = 1 (initial) + max_retries, with delays of 2, 4, 8 seconds
    between retries.

    Args:
        invoice_id: ID of the invoice to process
        max_retries: Maximum number of retries after the initial attempt (default: 3,
            resulting in up to 4 total attempts)
    """
    from app.services.invoice_service import process_invoice as do_process
    from app.database import async_session_maker
    from app.services.audit_service import log_audit
    import logging
    import asyncio
    logger = logging.getLogger(__name__)

    retry_count = 0
    last_error = None

    while retry_count <= max_retries:
        async with async_session_maker() as db:
            try:
                # Update status to PROCESSING
                query = select(Invoice).where(Invoice.id == invoice_id)
                result = await db.execute(query)
                invoice = result.scalar_one_or_none()

                if not invoice:
                    logger.error(f"Invoice {invoice_id} not found")
                    return

                invoice.status = InvoiceStatus.PROCESSING
                await db.commit()

                logger.info(f"Background processing invoice {invoice_id} (attempt {retry_count + 1}/{max_retries + 1})")
                success = await do_process(invoice_id, db)

                if success:
                    logger.info(f"Invoice {invoice_id} processing completed successfully")
                    # Log successful processing
                    await log_audit(
                        db=db,
                        entity_type="invoice",
                        entity_id=invoice_id,
                        action="process_complete",
                        new_value={"status": "success", "attempts": retry_count + 1}
                    )
                    return
                else:
                    last_error = "Processing returned false"
                    logger.warning(f"Invoice {invoice_id} processing returned false")

            except Exception as e:
                last_error = str(e)
                logger.error(f"Failed to process invoice {invoice_id} (attempt {retry_count + 1}): {e}")
                await db.rollback()

        retry_count += 1

        if retry_count <= max_retries:
            # Exponential backoff: 2^retry_count seconds (2, 4, 8 seconds)
            delay = 2 ** retry_count
            logger.info(f"Retrying invoice {invoice_id} in {delay} seconds...")
            await asyncio.sleep(delay)

    # All retries exhausted - mark as failed
    async with async_session_maker() as db:
        try:
            query = select(Invoice).where(Invoice.id == invoice_id)
            result = await db.execute(query)
            invoice = result.scalar_one_or_none()

            if invoice:
                # Set status back to UPLOADED so user can retry manually
                invoice.status = InvoiceStatus.UPLOADED
                await db.commit()

                # Log failed processing
                await log_audit(
                    db=db,
                    entity_type="invoice",
                    entity_id=invoice_id,
                    action="process_failed",
                    new_value={"error": last_error, "attempts": max_retries + 1}
                )

            logger.error(f"Invoice {invoice_id} processing failed after {max_retries + 1} attempts: {last_error}")
        except Exception as e:
            logger.error(f"Failed to update invoice {invoice_id} status after retry exhaustion: {e}")


@router.get("", response_model=InvoiceListResponse)
async def list_invoices(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(20, ge=1, le=100, description="每页数量"),
    status: Optional[InvoiceStatus] = Query(None, description="状态筛选"),
    owner: Optional[str] = Query(None, description="归属人筛选"),
    start_date: Optional[str] = Query(None, description="开始日期"),
    end_date: Optional[str] = Query(None, description="结束日期"),
    db: AsyncSession = Depends(get_db)
):
    """获取发票列表"""
    query = select(Invoice)

    # Apply filters
    if status:
        query = query.where(Invoice.status == status)
    if owner:
        query = query.where(Invoice.owner == owner)
    if start_date:
        try:
            start_date_obj = date.fromisoformat(start_date)
            query = query.where(Invoice.issue_date >= start_date_obj)
        except ValueError:
            pass  # Invalid date format, skip filter
    if end_date:
        try:
            end_date_obj = date.fromisoformat(end_date)
            query = query.where(Invoice.issue_date <= end_date_obj)
        except ValueError:
            pass  # Invalid date format, skip filter

    # Count total
    count_query = select(func.count()).select_from(query.subquery())
    total = await db.scalar(count_query) or 0

    # Apply pagination
    query = query.order_by(Invoice.created_at.desc())
    query = query.offset((page - 1) * page_size).limit(page_size)

    result = await db.execute(query)
    invoices = result.scalars().all()

    return InvoiceListResponse(
        items=[InvoiceResponse.model_validate(inv) for inv in invoices],
        total=total,
        page=page,
        page_size=page_size
    )


@router.get("/statistics", response_model=StatisticsResponse)
async def get_statistics(
    invoice_ids: Optional[str] = Query(None, description="发票ID列表，逗号分隔"),
    status: Optional[InvoiceStatus] = Query(None, description="状态筛选"),
    owner: Optional[str] = Query(None, description="归属人筛选"),
    db: AsyncSession = Depends(get_db)
):
    """获取发票统计数据"""
    query = select(Invoice)

    ids = _parse_invoice_ids(invoice_ids)
    if ids:
        query = query.where(Invoice.id.in_(ids))
    if status:
        query = query.where(Invoice.status == status)
    if owner:
        query = query.where(Invoice.owner == owner)

    result = await db.execute(query)
    invoices = result.scalars().all()

    count = len(invoices)
    total_amount = sum((inv.amount or Decimal(0)) for inv in invoices)
    total_tax = sum((inv.tax_amount or Decimal(0)) for inv in invoices)
    total_with_tax = sum((inv.total_with_tax or Decimal(0)) for inv in invoices)

    return StatisticsResponse(
        count=count,
        total_amount=total_amount,
        total_tax=total_tax,
        total_with_tax=total_with_tax
    )


@router.get("/{invoice_id}", response_model=InvoiceDetailResponse)
async def get_invoice(
    invoice_id: int,
    db: AsyncSession = Depends(get_db)
):
    """获取发票详情"""
    invoice = await _get_invoice_or_404(db, invoice_id)
    return await _build_detail(db, invoice)


@router.post("/{invoice_id}/retry-qr", response_model=QrRetryResponse)
async def retry_qr(
    invoice_id: int,
    db: AsyncSession = Depends(get_db)
):
    """手动触发二维码解码，并用解出的字段覆盖对应项。

    二维码线路不进默认链路：它在当前测试集上 51/51 解码失败
    （zbar 与 cv2 两个解码器、四档缩放、灰度变体全试过）。做成手动按钮，
    避免每张票都白跑一次；解出来时它是确定性来源，优先级高于任何模型。
    """
    invoice = await _get_invoice_or_404(db, invoice_id, for_update=True)
    _require_editable(invoice)

    from app.services.qr_service import decode_invoice_qr

    qr_fields = await run_in_threadpool(decode_invoice_qr, invoice.file_data)

    if not qr_fields:
        return QrRetryResponse(
            invoice_id=invoice_id, decoded=False, fields={},
            message="未能从图像中解出二维码",
        )

    # 二维码是确定性来源，直接覆盖并把对应行标为已确认。人填写过的字段不覆盖。
    diffs = await _get_diffs(db, invoice_id)
    by_name = {d.field_name: d for d in diffs}

    applied = {}
    changed = False
    for fname, value in qr_fields.items():
        if value is None:
            continue
        d = by_name.get(fname)
        if d is not None and d.source == SRC_MANUAL:
            continue
        try:
            # 按列类型转换后再写：日期、金额列不能直接塞字符串
            apply_to_invoice(invoice, fname, value)
        except InvalidValue:
            continue
        applied[fname] = value
        if d:
            changed = changed or d.final_value != value
            d.qr_value = value
            d.final_value = value
            d.source = "qr"
            d.confidence = "certain"
            d.note = "二维码解出（手动重试）"
            d.resolved = 1

    # 金额被二维码改写：整组按三层规则重判（二维码的值不会被改写）
    if "amount" in applied:
        for fname in reconcile_after_edit(by_name, "amount", app_settings.current().amount_rules):
            apply_to_invoice(invoice, fname, by_name[fname].final_value)

    remaining = [f for f in (invoice.degraded or []) if f not in applied]
    invoice.degraded = remaining
    _refresh_verdict(invoice, diffs)
    if changed and invoice.status == InvoiceStatus.CONFIRMED:
        invoice.status = InvoiceStatus.REVIEWING
    await db.commit()

    return QrRetryResponse(
        invoice_id=invoice_id, decoded=True, fields=applied,
        message=f"解码成功，已覆盖 {len(applied)} 个字段",
    )


@router.get("/{invoice_id}/file")
async def get_invoice_file(
    invoice_id: int,
    db: AsyncSession = Depends(get_db)
):
    """获取发票原始文件"""
    from fastapi.responses import Response
    from urllib.parse import quote

    query = select(Invoice).where(Invoice.id == invoice_id)
    result = await db.execute(query)
    invoice = result.scalar_one_or_none()

    if not invoice:
        raise HTTPException(status_code=404, detail="发票不存在")

    content_type_map = {
        "pdf": "application/pdf",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "png": "image/png"
    }

    # URL-encode filename for Content-Disposition header (RFC 5987)
    encoded_filename = quote(invoice.file_name)

    return Response(
        content=invoice.file_data,
        media_type=content_type_map.get(invoice.file_type, "application/octet-stream"),
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )


@router.get("/{invoice_id}/preview")
async def get_invoice_preview(
    invoice_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db)
):
    """原件预览：图片按 EXIF 转正、PDF 渲染首页，统一为 JPEG，供详情页查看器使用。

    /file 返回的是原文件且带 attachment，放进 iframe / img 会触发下载而不是显示。
    """
    invoice = await _get_invoice_or_404(db, invoice_id)

    # 原件入库后不会变，按内容做 ETag，重复打开同一张票不必重新渲染
    etag = f'"{PREVIEW_VERSION}-{hashlib.sha1(invoice.file_data).hexdigest()}"'
    headers = {"ETag": etag, "Cache-Control": "private, max-age=86400"}
    if request.headers.get("if-none-match") == etag:
        return Response(status_code=304, headers=headers)

    from app.services.preview_service import render_preview

    try:
        content = await run_in_threadpool(render_preview, invoice.file_data, invoice.file_type)
    except Exception as e:  # noqa: BLE001
        logger.warning("Invoice %s preview failed: %s: %s", invoice_id, type(e).__name__, e)
        raise HTTPException(status_code=500, detail="原件预览生成失败")

    return Response(
        content=content,
        media_type="image/jpeg",
        headers={**headers, "Content-Disposition": "inline"},
    )


@router.put("/{invoice_id}", response_model=InvoiceResponse)
async def update_invoice(
    invoice_id: int,
    update_data: InvoiceUpdate,
    request: Request,
    db: AsyncSession = Depends(get_db)
):
    """更新发票信息"""
    query = select(Invoice).where(Invoice.id == invoice_id)
    result = await db.execute(query)
    invoice = result.scalar_one_or_none()

    if not invoice:
        raise HTTPException(status_code=404, detail="发票不存在")

    # Capture old values for audit
    update_dict = update_data.model_dump(exclude_unset=True)
    old_values = {key: getattr(invoice, key) for key in update_dict.keys()}
    # Convert non-serializable types
    for key, value in old_values.items():
        if hasattr(value, 'value'):  # Enum
            old_values[key] = value.value
        elif hasattr(value, 'isoformat'):  # Date/DateTime
            old_values[key] = value.isoformat()
        elif isinstance(value, Decimal):
            old_values[key] = str(value)

    for key, value in update_dict.items():
        setattr(invoice, key, value)

    # Audit log
    client_info = get_client_info(request)
    new_values = update_dict.copy()
    for key, value in new_values.items():
        if hasattr(value, 'value'):  # Enum
            new_values[key] = value.value
        elif hasattr(value, 'isoformat'):  # Date/DateTime
            new_values[key] = value.isoformat()
        elif isinstance(value, Decimal):
            new_values[key] = str(value)

    await log_audit_no_commit(
        db=db,
        entity_type="invoice",
        entity_id=invoice_id,
        action="update",
        old_value=old_values,
        new_value=new_values,
        ip_address=client_info.get("ip_address"),
        user_agent=client_info.get("user_agent"),
    )

    await db.commit()
    await db.refresh(invoice)

    return InvoiceResponse.model_validate(invoice)


@router.post("/batch-update")
@limiter.limit("30/minute")
async def batch_update_invoices(
    request: Request,
    batch_request: BatchUpdateRequest,
    db: AsyncSession = Depends(get_db)
):
    """批量更新发票状态/归属人"""
    query = select(Invoice).where(Invoice.id.in_(batch_request.invoice_ids))
    result = await db.execute(query)
    invoices = result.scalars().all()

    client_info = get_client_info(request)
    updated_count = 0
    for invoice in invoices:
        old_values = {}
        new_values = {}
        if batch_request.status is not None:
            old_values["status"] = invoice.status.value if invoice.status else None
            invoice.status = batch_request.status
            new_values["status"] = batch_request.status.value
        if batch_request.owner is not None:
            old_values["owner"] = invoice.owner
            invoice.owner = batch_request.owner
            new_values["owner"] = batch_request.owner
        updated_count += 1

        # Audit log for each invoice
        await log_audit_no_commit(
            db=db,
            entity_type="invoice",
            entity_id=invoice.id,
            action="batch_update",
            old_value=old_values,
            new_value=new_values,
            ip_address=client_info.get("ip_address"),
            user_agent=client_info.get("user_agent"),
        )

    await db.commit()

    return {
        "message": f"成功更新 {updated_count} 张发票",
        "updated_count": updated_count
    }


@router.post("/batch-delete")
@limiter.limit("20/minute")
async def batch_delete_invoices(
    request: Request,
    batch_request: BatchDeleteRequest,
    db: AsyncSession = Depends(get_db)
):
    """批量删除发票及其关联数据"""
    if not batch_request.invoice_ids:
        raise HTTPException(status_code=400, detail="请选择要删除的发票")

    # Query invoices to delete
    query = select(Invoice).where(Invoice.id.in_(batch_request.invoice_ids))
    result = await db.execute(query)
    invoices = result.scalars().all()

    if not invoices:
        raise HTTPException(status_code=404, detail="未找到要删除的发票")

    client_info = get_client_info(request)
    deleted_count = 0
    for invoice in invoices:
        # Audit log for each deletion
        await log_audit_no_commit(
            db=db,
            entity_type="invoice",
            entity_id=invoice.id,
            action="delete",
            old_value={"file_name": invoice.file_name, "invoice_number": invoice.invoice_number},
            ip_address=client_info.get("ip_address"),
            user_agent=client_info.get("user_agent"),
        )
        await db.delete(invoice)
        deleted_count += 1

    await db.commit()

    return {
        "message": f"成功删除 {deleted_count} 张发票",
        "deleted_count": deleted_count
    }


@router.post("/batch-reprocess")
@limiter.limit("5/minute")
async def batch_reprocess_invoices(
    request: Request,
    background_tasks: BackgroundTasks,
    batch_request: BatchDeleteRequest,  # Reuse for invoice_ids
    db: AsyncSession = Depends(get_db)
):
    """批量重新解析发票（清除旧的OCR/LLM结果，重新处理）"""
    import logging
    logger = logging.getLogger(__name__)

    if not batch_request.invoice_ids:
        raise HTTPException(status_code=400, detail="请选择要重新解析的发票")

    # Query invoices to reprocess
    query = select(Invoice).where(Invoice.id.in_(batch_request.invoice_ids))
    result = await db.execute(query)
    invoices = result.scalars().all()

    if not invoices:
        raise HTTPException(status_code=404, detail="未找到要重新解析的发票")

    # Clear old parsing results and reset invoice fields
    for invoice in invoices:
        # Delete old OCR results
        ocr_query = select(OcrResult).where(OcrResult.invoice_id == invoice.id)
        ocr_result = await db.execute(ocr_query)
        for ocr in ocr_result.scalars().all():
            await db.delete(ocr)

        # Delete old LLM results
        llm_query = select(LlmResult).where(LlmResult.invoice_id == invoice.id)
        llm_result = await db.execute(llm_query)
        for llm in llm_result.scalars().all():
            await db.delete(llm)

        # Delete old parsing diffs
        diff_query = select(ParsingDiff).where(ParsingDiff.invoice_id == invoice.id)
        diff_result = await db.execute(diff_query)
        for diff in diff_result.scalars().all():
            await db.delete(diff)

        # Reset invoice fields
        invoice.invoice_number = None
        invoice.issue_date = None
        invoice.buyer_name = None
        invoice.buyer_tax_id = None
        invoice.seller_name = None
        invoice.seller_tax_id = None
        invoice.item_name = None
        invoice.total_with_tax = None
        invoice.amount = None
        invoice.tax_amount = None
        invoice.tax_rate = None
        invoice.status = InvoiceStatus.UPLOADED

    await db.commit()
    logger.info(f"Cleared old parsing results for {len(invoices)} invoices, scheduling reprocess")

    # Schedule background processing
    for invoice in invoices:
        background_tasks.add_task(process_invoice_background, invoice.id)

    return {
        "message": f"已清除 {len(invoices)} 张发票的旧解析结果，正在重新解析",
        "count": len(invoices)
    }


@router.post("/{invoice_id}/process")
async def process_invoice(
    invoice_id: int,
    db: AsyncSession = Depends(get_db)
):
    """处理发票：运行OCR解析"""
    from app.services.invoice_service import process_invoice as do_process

    invoice = await _get_invoice_or_404(db, invoice_id)
    # 已上传 / 解析中的票也允许手动（重新）识别：后台任务可能已经失败或随进程重启丢失
    _require_editable(invoice, allow_unprocessed=True)

    success = await do_process(invoice_id, db)

    if success:
        return {"message": "解析成功", "invoice_id": invoice_id}
    else:
        raise HTTPException(status_code=500, detail="解析失败")


@router.delete("/{invoice_id}")
async def delete_invoice(
    invoice_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db)
):
    """删除发票"""
    query = select(Invoice).where(Invoice.id == invoice_id)
    result = await db.execute(query)
    invoice = result.scalar_one_or_none()

    if not invoice:
        raise HTTPException(status_code=404, detail="发票不存在")

    # Audit log
    client_info = get_client_info(request)
    await log_audit_no_commit(
        db=db,
        entity_type="invoice",
        entity_id=invoice_id,
        action="delete",
        old_value={"file_name": invoice.file_name, "invoice_number": invoice.invoice_number},
        ip_address=client_info.get("ip_address"),
        user_agent=client_info.get("user_agent"),
    )

    await db.delete(invoice)
    await db.commit()

    return {"message": "删除成功"}


@router.patch("/{invoice_id}/fields/{field_name}", response_model=InvoiceDetailResponse)
async def update_invoice_field(
    invoice_id: int,
    field_name: str,
    body: FieldUpdateRequest,
    request: Request,
    db: AsyncSession = Depends(get_db)
):
    """详情页对单个字段的确认或修改。

    所有字段操作都走这里（选中某个候选值也是 set）：写 ParsingDiff 行的同时写
    invoices 列，金额组（金额、税率、税额、价税合计）任一项被确认或修改后整组重判，
    最后按当前值重校验。
    返回完整详情，前端直接替换状态，不必再请求一次。
    """
    invoice = await _get_invoice_or_404(db, invoice_id, for_update=True)
    _require_editable(invoice)
    client_info = get_client_info(request)

    if field_name == "owner":
        # 归属人不是识别结果，没有 ParsingDiff 行：只写列，不参与校验，也不影响状态
        if body.action != "set":
            raise HTTPException(status_code=400, detail="归属人只支持修改")
        value = (body.value or "").strip() or None
        if value and len(value) > 100:
            raise HTTPException(status_code=400, detail="内容过长，最多 100 个字符")
        old_owner = invoice.owner
        invoice.owner = value
        await log_audit_no_commit(
            db=db,
            entity_type="invoice",
            entity_id=invoice_id,
            action="field_update",
            old_value={"field_name": "owner", "value": old_owner},
            new_value={"field_name": "owner", "value": value},
            ip_address=client_info.get("ip_address"),
            user_agent=client_info.get("user_agent"),
        )
        await db.commit()
        return await _build_detail(db, invoice)

    if field_name not in DIFF_FIELDS:
        raise HTTPException(status_code=404, detail="不支持修改该字段")
    diffs = await _get_diffs(db, invoice_id)
    by_name = {d.field_name: d for d in diffs}
    row = by_name.get(field_name)
    if row is None:
        raise HTTPException(status_code=404, detail="这张发票没有该字段的识别结果")

    before = {"final_value": row.final_value, "source": row.source, "resolved": row.resolved}
    touched: List[str] = []
    rules = app_settings.current().amount_rules
    try:
        if body.action == "confirm":
            if row.final_value is None and field_name in REQUIRED_FIELDS:
                raise InvalidValue("这一项必填，请先填写")
            row.resolved = 1
        elif body.action == "set":
            value, source, same = resolve_set(row, body.value)
            if not same:
                # 先写列：类型转换失败时 diff 行还没被改动
                apply_to_invoice(invoice, field_name, value)
                row.final_value = value
                row.source = source
                row.note = "人工填写"
                touched.append(field_name)
            row.resolved = 1
        else:  # auto：金额组某一项交还给公式，按其余三项推导
            touched = restore_auto(by_name, field_name, rules)
        if body.action != "auto" and field_name in AMOUNT_GROUP:
            touched += [f for f in reconcile_after_edit(by_name, field_name, rules)
                        if f not in touched]
        for fname in touched:
            apply_to_invoice(invoice, fname, by_name[fname].final_value)
    except InvalidValue as e:
        raise HTTPException(status_code=400, detail=str(e))

    _refresh_verdict(invoice, diffs)
    if touched and invoice.status == InvoiceStatus.CONFIRMED:
        # 改了已确认发票的值，要重新确认
        invoice.status = InvoiceStatus.REVIEWING

    await log_audit_no_commit(
        db=db,
        entity_type="parsing_diff",
        entity_id=row.id,
        action="field_update",
        old_value={"field_name": field_name, **before},
        new_value={
            "field_name": field_name,
            "action": body.action,
            "final_value": row.final_value,
            "source": row.source,
            "resolved": row.resolved,
            "recomputed": [f for f in touched if f != field_name],
            "invoice_id": invoice_id,
        },
        ip_address=client_info.get("ip_address"),
        user_agent=client_info.get("user_agent"),
    )

    await db.commit()
    return await _build_detail(db, invoice)


@router.post("/{invoice_id}/confirm")
async def confirm_invoice(
    invoice_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db)
):
    """确认发票。还没核对的字段按当前结果一并确认（前端确认前会列出这些字段）。"""
    invoice = await _get_invoice_or_404(db, invoice_id, for_update=True)
    _require_editable(invoice)

    missing_fields = [field for field in REQUIRED_FIELDS if _is_blank(getattr(invoice, field))]
    if missing_fields:
        raise HTTPException(
            status_code=400,
            detail="无法确认：必填字段缺失，请补全后再确认。"
        )

    # Check if LLM result exists
    llm_query = select(LlmResult).where(LlmResult.invoice_id == invoice_id)
    llm_result = await db.execute(llm_query)
    llm = llm_result.scalar_one_or_none()
    has_llm = llm is not None

    diffs = await _get_diffs(db, invoice_id)

    if has_llm and not diffs:
        raise HTTPException(
            status_code=400,
            detail="无法确认：缺少OCR和LLM的比对结果。请重新解析发票。"
        )

    # 销售方与本公司主体相同是硬矛盾，不能靠「一并确认」带过去
    _refresh_verdict(invoice, diffs)
    if invoice.pipeline_version == PIPELINE_VERSION and invoice.verdict == VERDICT_REJECTED:
        raise HTTPException(
            status_code=400,
            detail="无法确认：销售方与本公司主体相同，请核对销售方后再确认。"
        )

    # 其余未核对的字段按当前结果一并确认。不再用某一路识别值去补空值——
    # 那会把被 seller_is_entity 置空的值又填回去。
    auto_accepted = [d.field_name for d in diffs if not d.resolved]
    for diff in diffs:
        diff.resolved = 1
    _refresh_verdict(invoice, diffs)

    # Update invoice status
    old_status = invoice.status.value if invoice.status else None
    invoice.status = InvoiceStatus.CONFIRMED

    # Audit log for confirmation
    client_info = get_client_info(request)
    await log_audit_no_commit(
        db=db,
        entity_type="invoice",
        entity_id=invoice_id,
        action="confirm",
        old_value={"status": old_status},
        new_value={
            "status": InvoiceStatus.CONFIRMED.value,
            "resolved_diffs": len(diffs),
            "auto_accepted": auto_accepted,
            "source": "llm" if has_llm else "ocr",
        },
        ip_address=client_info.get("ip_address"),
        user_agent=client_info.get("user_agent"),
    )

    await db.commit()

    return {"message": "发票已确认", "resolved_count": len(diffs)}


@router.get("/export/csv")
@limiter.limit("10/minute")
async def export_invoices_csv(
    request: Request,
    invoice_ids: Optional[str] = Query(None, description="发票ID列表，逗号分隔"),
    status: Optional[InvoiceStatus] = Query(None, description="状态筛选"),
    owner: Optional[str] = Query(None, description="归属人筛选"),
    start_date: Optional[str] = Query(None, description="开始日期"),
    end_date: Optional[str] = Query(None, description="结束日期"),
    db: AsyncSession = Depends(get_db)
):
    """导出发票为CSV格式"""
    import csv
    from urllib.parse import quote

    query = select(Invoice)

    ids = _parse_invoice_ids(invoice_ids)
    if ids:
        query = query.where(Invoice.id.in_(ids))
    if status:
        query = query.where(Invoice.status == status)
    if owner:
        query = query.where(Invoice.owner == owner)
    start_date_obj = _parse_date_param(start_date, "start_date")
    if start_date_obj:
        query = query.where(Invoice.issue_date >= start_date_obj)
    end_date_obj = _parse_date_param(end_date, "end_date")
    if end_date_obj:
        query = query.where(Invoice.issue_date <= end_date_obj)

    query = query.order_by(Invoice.created_at.desc())
    result = await db.execute(query)
    invoices = result.scalars().all()

    # Create CSV content
    output = BytesIO()
    # Write BOM for Excel compatibility with Chinese characters
    output.write(b'\xef\xbb\xbf')

    import codecs
    writer = csv.writer(codecs.getwriter('utf-8')(output))

    # Header row
    writer.writerow([
        '发票号码', '开票日期', '购买方名称', '购买方纳税人识别号',
        '销售方名称', '销售方纳税人识别号', '项目名称',
        '金额', '税额', '价税合计', '税率',
        '状态', '归属人', '文件名', '创建时间'
    ])

    # Data rows
    for inv in invoices:
        writer.writerow([
            inv.invoice_number or '',
            str(inv.issue_date) if inv.issue_date else '',
            inv.buyer_name or '',
            inv.buyer_tax_id or '',
            inv.seller_name or '',
            inv.seller_tax_id or '',
            inv.item_name or '',
            str(inv.amount) if inv.amount else '',
            str(inv.tax_amount) if inv.tax_amount else '',
            str(inv.total_with_tax) if inv.total_with_tax else '',
            inv.tax_rate or '',
            inv.status.value if inv.status else '',
            inv.owner or '',
            inv.file_name or '',
            inv.created_at.strftime('%Y-%m-%d %H:%M:%S') if inv.created_at else ''
        ])

    output.seek(0)

    filename = quote('发票导出.csv')
    return StreamingResponse(
        output,
        media_type='text/csv; charset=utf-8',
        headers={
            'Content-Disposition': f"attachment; filename*=UTF-8''{filename}"
        }
    )


@router.get("/export/excel")
@limiter.limit("10/minute")
async def export_invoices_excel(
    request: Request,
    invoice_ids: Optional[str] = Query(None, description="发票ID列表，逗号分隔"),
    status: Optional[InvoiceStatus] = Query(None, description="状态筛选"),
    owner: Optional[str] = Query(None, description="归属人筛选"),
    start_date: Optional[str] = Query(None, description="开始日期"),
    end_date: Optional[str] = Query(None, description="结束日期"),
    db: AsyncSession = Depends(get_db)
):
    """导出发票为Excel格式"""
    from urllib.parse import quote

    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="Excel导出需要安装openpyxl库")

    query = select(Invoice)

    ids = _parse_invoice_ids(invoice_ids)
    if ids:
        query = query.where(Invoice.id.in_(ids))
    if status:
        query = query.where(Invoice.status == status)
    if owner:
        query = query.where(Invoice.owner == owner)
    start_date_obj = _parse_date_param(start_date, "start_date")
    if start_date_obj:
        query = query.where(Invoice.issue_date >= start_date_obj)
    end_date_obj = _parse_date_param(end_date, "end_date")
    if end_date_obj:
        query = query.where(Invoice.issue_date <= end_date_obj)

    query = query.order_by(Invoice.created_at.desc())
    result = await db.execute(query)
    invoices = result.scalars().all()

    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "发票列表"

    # Header row
    headers = [
        '发票号码', '开票日期', '购买方名称', '购买方纳税人识别号',
        '销售方名称', '销售方纳税人识别号', '项目名称',
        '金额', '税额', '价税合计', '税率',
        '状态', '归属人', '文件名', '创建时间'
    ]

    header_font = Font(bold=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')

    # Data rows
    for row, inv in enumerate(invoices, 2):
        data = [
            inv.invoice_number or '',
            str(inv.issue_date) if inv.issue_date else '',
            inv.buyer_name or '',
            inv.buyer_tax_id or '',
            inv.seller_name or '',
            inv.seller_tax_id or '',
            inv.item_name or '',
            float(inv.amount) if inv.amount else '',
            float(inv.tax_amount) if inv.tax_amount else '',
            float(inv.total_with_tax) if inv.total_with_tax else '',
            inv.tax_rate or '',
            inv.status.value if inv.status else '',
            inv.owner or '',
            inv.file_name or '',
            inv.created_at.strftime('%Y-%m-%d %H:%M:%S') if inv.created_at else ''
        ]

        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = quote('发票导出.xlsx')
    return StreamingResponse(
        output,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={
            'Content-Disposition': f"attachment; filename*=UTF-8''{filename}"
        }
    )
